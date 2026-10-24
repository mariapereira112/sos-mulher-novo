# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE   = '/home/runner/workspace/attached_assets'
ASSETS = {
    'uc'     : f'{BASE}/uc_bw.png',
    'classes': f'{BASE}/classes_bw.png',
    'mvc'    : f'{BASE}/mvc_bw.png',
    'estrut' : f'{BASE}/estrutura2_bw.png',
    'proto'  : f'{BASE}/prototipacao.jpg',
    'twilio' : f'{BASE}/API-wtilio_1782252099065.jpg',
}

CRIM  = RGBColor(230, 57,  70)
WHITE = RGBColor(255,255,255)
DARK  = RGBColor( 28, 28, 30)
LGRAY = RGBColor(245,245,245)
BLACK = RGBColor(  0,  0,  0)

def hex_color(r,g,b):
    return f'{r:02X}{g:02X}{b:02X}'

def set_cell_bg(cell, r,g,b):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color(r,g,b))
    tcPr.append(shd)

def add_heading(doc, text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.color.rgb = color or CRIM
    return p

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(10)
    return p

def add_bullet(doc, items):
    for item in items:
        p = doc.add_paragraph(f'• {item}', style='Normal')
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(2)

def add_figure(doc, path, caption, width=14):
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90,90,90)
    doc.add_paragraph()

def add_req_table(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, (hdr, w) in enumerate(zip(headers, col_widths)):
        hdr_cells[i].text = hdr
        hdr_cells[i].width = Cm(w)
        set_cell_bg(hdr_cells[i], 230,57,70)
        for run in hdr_cells[i].paragraphs[0].runs:
            run.font.color.rgb = WHITE
            run.bold = True
            run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, (val, w) in enumerate(zip(row, col_widths)):
            cells[i].text = str(val)
            cells[i].width = Cm(w)
            if i == 0:
                set_cell_bg(cells[i], 245,245,245)
                for run in cells[i].paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
            else:
                for run in cells[i].paragraphs[0].runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()

# ── Document ──────────────────────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)

# ── CAPA ──────────────────────────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(60)
run = title.add_run('SOS MULHER')
run.font.size = Pt(32)
run.bold = True
run.font.color.rgb = CRIM

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Aplicativo de Seguranca para Mulheres em Situacao de Vulnerabilidade')
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(80,80,80)

doc.add_paragraph()

meta_table = doc.add_table(rows=5, cols=2)
meta_table.style = 'Table Grid'
metas = [
    ('Disciplina',    'Projeto Integrador 2 – Avaliacao 2'),
    ('Equipe',        'Maria Janaina'),
    ('Tecnologia',    'Node.js 24 | TypeScript 5.9 | Expo React Native | PostgreSQL (Neon Cloud)'),
    ('Notificacoes',  'Twilio WhatsApp + Voz Automatizada'),
    ('Data',          'Junho de 2026'),
]
for i,(lbl,val) in enumerate(metas):
    meta_table.rows[i].cells[0].text = lbl
    meta_table.rows[i].cells[1].text = val
    set_cell_bg(meta_table.rows[i].cells[0], 230,57,70)
    for run in meta_table.rows[i].cells[0].paragraphs[0].runs:
        run.font.color.rgb = WHITE
        run.bold = True
        run.font.size = Pt(10)
    for run in meta_table.rows[i].cells[1].paragraphs[0].runs:
        run.font.size = Pt(10)

doc.add_page_break()

# ── 1. INTRODUCAO ─────────────────────────────────────────────────────────────
add_heading(doc, '1. Introducao e Contextualizacao')
add_body(doc,
    'O SOS Mulher e um aplicativo movel premium desenvolvido para oferecer seguranca imediata '
    'a mulheres em situacoes de vulnerabilidade ou perigo. O sistema permite acionar alertas '
    'de emergencia com um unico toque prolongado, notificando automaticamente contatos de '
    'confianca via WhatsApp e ligacao de voz automatizada.')
add_body(doc,
    'O projeto foi concebido como resposta a realidade brasileira de violencia de genero, '
    'integrando tecnologias modernas de comunicacao (Twilio), banco de dados em nuvem '
    '(PostgreSQL Neon Cloud) e interface mobile nativa (Expo React Native).')

add_heading(doc, 'Objetivo Geral', level=2)
add_body(doc,
    'Desenvolver uma solucao tecnologica full-stack que permita a mulheres acionarem alertas '
    'de emergencia georreferenciados, com escalonamento automatico de notificacoes e '
    'registro auditavel de ocorrencias.')

add_heading(doc, 'Objetivos Especificos', level=2)
add_bullet(doc, [
    'Implementar autenticacao segura com hash de senha e recuperacao de conta por e-mail.',
    'Criar fluxo de SOS: 1,5 s de pressionamento → WhatsApp → Ligacao de voz automatizada.',
    'Persistir historico de alertas e status no PostgreSQL Neon Cloud.',
    'Oferecer comunidade, boletins de ocorrencia e gerenciamento de contatos.',
    'Garantir acesso somente a mulheres e pessoas nao-binarias.',
    'Produzir codigo limpo, tipado (TypeScript) seguindo padrao MVC.',
])

add_heading(doc, 'Stack Tecnologica', level=2)
stack_rows = [
    ('Mobile','Expo 52 + React Native + Expo Router'),
    ('Backend','Node.js 24 + Express 5 + TypeScript 5.9'),
    ('Banco de Dados','PostgreSQL 16 – Neon Cloud (serverless, SSL)'),
    ('ORM','Drizzle ORM + drizzle-kit'),
    ('Validacao','Zod v4 + Orval codegen (OpenAPI 3.1)'),
    ('Notificacoes','Twilio WhatsApp Sandbox + Twilio Voice'),
    ('Monorepo','pnpm workspaces (db, api-spec, api-client-react, api-zod)'),
]
add_req_table(doc, ['Camada','Descricao'], stack_rows, [5, 12])

doc.add_page_break()

# ── 2. REQUISITOS ─────────────────────────────────────────────────────────────
add_heading(doc, '2. Requisitos do Sistema')
add_heading(doc, '2.1 Requisitos Funcionais (RF)', level=2)
rfs = [
    ('RF01','Cadastrar usuario com nome, CPF, e-mail, senha e genero','Alta'),
    ('RF02','Bloquear cadastro para genero "Homem" – redirect para tela restrita','Alta'),
    ('RF03','Realizar login por nome + senha com hash bcrypt','Alta'),
    ('RF04','Recuperar conta por e-mail (token simulado)','Media'),
    ('RF05','Acionar alerta SOS com pressionamento de 1,5 segundo','Alta'),
    ('RF06','Enviar WhatsApp para todos os contatos ao criar alerta (Fase 1)','Alta'),
    ('RF07','Iniciar ligacao de voz automatizada apos temporizador (Fase 2)','Alta'),
    ('RF08','Resolver alerta via webhook ao pressionar digito 1','Alta'),
    ('RF09','Cancelar alerta ativo com protecao contra race condition','Alta'),
    ('RF10','Registrar historico de alertas com status e timestamp','Alta'),
    ('RF11','Gerenciar ate 5 contatos de emergencia personalizados','Alta'),
    ('RF12','Exibir linhas oficiais de ajuda (180, 190, 192)','Alta'),
    ('RF13','Publicar e visualizar posts na comunidade','Media'),
    ('RF14','Registrar boletim de ocorrencia com descricao e data','Media'),
    ('RF15','Persistir configuracoes de seguranca e temporizador no banco','Media'),
]
add_req_table(doc, ['ID','Descricao','Prioridade'], rfs, [2, 13, 3])

add_heading(doc, '2.2 Requisitos Nao Funcionais (RNF)', level=2)
rnfs = [
    ('RNF01','Banco de dados PostgreSQL 16 na Neon Cloud com SSL','Alta'),
    ('RNF02','API REST documentada em OpenAPI 3.1 (contract-first)','Alta'),
    ('RNF03','Toda comunicacao API via HTTPS/TLS','Alta'),
    ('RNF04','Senhas armazenadas com bcrypt (hash+salt)','Alta'),
    ('RNF05','Codegen automatico: OpenAPI → Zod + React Query hooks','Alta'),
    ('RNF06','Monorepo pnpm com pacotes compartilhados tipados','Media'),
    ('RNF07','Logs estruturados JSON via Pino (nunca console.log)','Media'),
    ('RNF08','App em tema escuro permanente (#1c1c1e / #e63946)','Media'),
    ('RNF09','Resposta da API < 500ms em operacoes nao-SOS','Media'),
    ('RNF10','TypeScript strict em todos os pacotes','Alta'),
]
add_req_table(doc, ['ID','Descricao','Prioridade'], rnfs, [2, 13, 3])

doc.add_page_break()

# ── 3. CASOS DE USO ──────────────────────────────────────────────────────────
add_heading(doc, '3. Casos de Uso')
add_body(doc,
    'O diagrama abaixo apresenta os principais casos de uso do sistema SOS Mulher, '
    'identificando atores externos e interacoes com o sistema.')
add_figure(doc, ASSETS['uc'], 'Figura 1 – Diagrama de Casos de Uso (UML)', width=13)

add_heading(doc, 'Descricao dos Casos de Uso', level=2)
ucs = [
    ('UC01 – Login','Usuario insere nome e senha; sistema valida hash e inicia sessao.'),
    ('UC02 – Cadastro','Coleta CPF, genero, e-mail, senha; bloqueia genero "Homem".'),
    ('UC03 – Acionar SOS','Pressionamento de 1,5 s cria alerta; inclui UC04 e estende UC05.'),
    ('UC04 – Enviar WhatsApp','Disparo imediato de mensagem Twilio para contatos cadastrados.'),
    ('UC05 – Ligacao de Voz','Apos temporizador, Twilio efetua chamada automatizada.'),
    ('UC06 – Cancelar Alerta','Cancela alerta ativo; guard impede chamadas em andamento.'),
    ('UC07 – Ver Historico','Lista alertas com status (pending/sms_sent/resolved).'),
    ('UC08 – Boletim de Ocorrencia','Registra descricao de incidente com data.'),
    ('UC09 – Gerenciar Contatos','CRUD de ate 5 contatos de emergencia personalizados.'),
    ('UC10 – Configurar Perfil','Atualiza dados pessoais, permissoes e temporizador.'),
]
add_req_table(doc, ['Caso de Uso','Descricao'], ucs, [5, 12])

doc.add_page_break()

# ── 4. ARQUITETURA ────────────────────────────────────────────────────────────
add_heading(doc, '4. Arquitetura do Sistema')
add_body(doc,
    'O sistema adota o padrao MVC (Model-View-Controller) com separacao clara '
    'de responsabilidades, implementado em TypeScript com Drizzle ORM para '
    'acesso ao banco de dados Neon Cloud.')

add_heading(doc, '4.1 Padrao MVC com TypeScript e Drizzle ORM', level=2)
add_figure(doc, ASSETS['mvc'], 'Figura 2 – Arquitetura MVC (TypeScript + Drizzle + Neon Cloud)', width=14)

layers = [
    ('Model (Drizzle ORM)','UserModel, AlertModel, FeedbackModel – queries tipadas ao PostgreSQL Neon Cloud.'),
    ('Service (Logica)','AlertService, AuthService – SOS escalation, hash senha, Twilio.'),
    ('Controller (HTTP)','alertController, authController – valida input (Zod), chama Service, retorna JSON.'),
    ('Route (Express 5)','/api/alerts, /api/auth, /api/users, /api/contacts, /api/settings, /api/posts.'),
    ('Codegen (Orval)','OpenAPI 3.1 → React Query hooks + Zod schemas (gerados automaticamente).'),
]
add_req_table(doc, ['Camada','Descricao'], layers, [5, 12])

add_heading(doc, '4.2 Diagrama de Classes', level=2)
add_figure(doc, ASSETS['classes'], 'Figura 3 – Diagrama de Classes (UML – Notacao Portuguesa)', width=14)

add_heading(doc, '4.3 Estrutura de Modelos (POO)', level=2)
add_body(doc,
    'Os modelos seguem orientacao a objetos com encapsulamento das queries Drizzle, '
    'retornando objetos tipados para as camadas superiores.')
add_figure(doc, ASSETS['estrut'], 'Figura 4 – Estrutura de Modelos POO (TypeScript)', width=13)

doc.add_page_break()

# ── 5. BANCO DE DADOS ────────────────────────────────────────────────────────
add_heading(doc, '5. Banco de Dados – PostgreSQL Neon Cloud')
add_body(doc,
    'O projeto utiliza PostgreSQL 16 hospedado na Neon Cloud, plataforma serverless '
    'que oferece escalabilidade automatica, backups gerenciados e conexao segura via SSL. '
    'O ORM Drizzle gerencia o schema e as migracoes.')
add_body(doc, 'String de conexao (formato Neon Cloud):')
p = doc.add_paragraph()
run = p.add_run('postgresql://usuario:senha@ep-xxxx.us-east-2.aws.neon.tech/sos_mulher?sslmode=require')
run.font.name = 'Courier New'
run.font.size = Pt(9)
run.font.color.rgb = CRIM

add_heading(doc, 'Tabelas do Schema', level=2)
tables = [
    ('users','id, name, email, cpf, gender, passwordHash, onboardingComplete, createdAt'),
    ('emergency_contacts','id, userId, name, phone, relationship, createdAt'),
    ('alerts','id, userId, status, latitude, longitude, timerDuration, voicePath, createdAt'),
    ('alert_notifications','id, alertId, contactId, channel, status, sentAt'),
    ('community_posts','id, userId, content, category, createdAt'),
    ('settings','id, userId, silentMode, disguiseMode, timerDuration, shareLocation, updatedAt'),
    ('incidents','id, userId, description, location, occurredAt, createdAt'),
    ('feedbacks','id, userId, type, message, createdAt'),
]
add_req_table(doc, ['Tabela','Colunas Principais'], tables, [5, 12])

add_heading(doc, 'Status do Alerta', level=2)
estados = [
    ('pending','Alerta criado, aguardando envio de notificacoes.'),
    ('sms_sent','WhatsApp enviados para todos os contatos (Fase 1 concluida).'),
    ('calling_cloud','Ligacao de voz iniciada via Twilio apos temporizador (Fase 2).'),
    ('resolved_by_contact','Contato pressionou digito 1 – emergencia resolvida.'),
    ('cancelled','Usuario cancelou; guard impede chamadas em andamento.'),
]
add_req_table(doc, ['Status','Descricao'], estados, [5, 12])

doc.add_page_break()

# ── 6. TWILIO ────────────────────────────────────────────────────────────────
add_heading(doc, '6. Integracao e Notificacoes – Twilio')
add_body(doc,
    'O sistema utiliza Twilio para WhatsApp (imediato) e ligacoes de voz automatizadas '
    '(Fase 2 do SOS). A integracao opera de forma graceful – se as variaveis nao '
    'estiverem presentes, o sistema simula o envio sem erros.')

add_heading(doc, 'Fluxo de Escalacao SOS (3 Fases)', level=2)
fases = [
    ('Fase 1 – WhatsApp','POST /api/alerts cria alerta + envia mensagem Twilio para cada contato. Status: sms_sent.'),
    ('Fase 2 – Ligacao de Voz','setTimeout dispara POST /api/alerts/trigger-voice. Guard verifica status antes de ligar. Status: calling_cloud.'),
    ('Fase 3 – Webhook','Contato pressiona 1 na chamada. PUT /api/alerts/:id/webhook/respond resolve. Status: resolved_by_contact.'),
]
add_req_table(doc, ['Fase','Descricao'], fases, [5, 12])

add_heading(doc, 'Comprovante Twilio WhatsApp', level=2)
add_figure(doc, ASSETS['twilio'], 'Figura 5 – Mensagem WhatsApp recebida via Twilio Sandbox', width=10)

add_heading(doc, 'Variaveis de Ambiente', level=2)
envvars = [
    ('TWILIO_ACCOUNT_SID','SID da conta Twilio'),
    ('TWILIO_AUTH_TOKEN','Token de autenticacao'),
    ('TWILIO_PHONE_NUMBER','Numero Twilio origem'),
    ('EMERGENCY_DESTINATION_NUMBER','Numero destino para voz'),
    ('SESSION_SECRET','Segredo para sessoes Express'),
    ('DATABASE_URL','String Neon Cloud com sslmode=require'),
]
add_req_table(doc, ['Variavel','Descricao'], envvars, [7, 10])

doc.add_page_break()

# ── 7. PROTOTIPACAO ──────────────────────────────────────────────────────────
add_heading(doc, '7. Prototipacao das Telas')
add_body(doc,
    'Telas implementadas com tema escuro permanente (#1c1c1e de fundo, #e63946 crimson). '
    'App utiliza Expo Router com 5 abas e botao SOS central elevado.')
add_figure(doc, ASSETS['proto'], 'Figura 6 – Prototipacao das Telas Principais (tema escuro)', width=14)

add_heading(doc, 'Telas Implementadas', level=2)
telas = [
    ('Onboarding (3 passos)','Intro → Aviso Legal → Cadastro com CPF + seletor de genero'),
    ('Acesso Restrito','Exibida ao genero "Homem" – sem acesso ao app'),
    ('Dashboard SOS','Botao SOS (1,5 s) + cards GPS/contatos + check-in'),
    ('Historico','Lista de alertas do DB com badges de status'),
    ('Comunidade','Feed cronologico de posts da base de dados'),
    ('Contatos','Linhas 180/190/192 + ate 5 contatos personalizados'),
    ('Perfil','4 accordions: Dados Pessoais, Permissoes, Seguranca, Temporizador'),
    ('Boletim de Ocorrencia','Formulario para registro de incidentes'),
    ('Modo Disfarce','Tela discreta para situacoes de perigo imediato'),
]
add_req_table(doc, ['Tela','Descricao'], telas, [5, 12])

doc.add_page_break()

# ── 8. CONCLUSAO ─────────────────────────────────────────────────────────────
add_heading(doc, '8. Conclusao e Consideracoes Finais')
add_body(doc,
    'O SOS Mulher representa uma solucao tecnologica completa e funcional para a seguranca '
    'de mulheres em situacoes de vulnerabilidade. O projeto entrega um sistema full-stack '
    'robusto com arquitetura bem definida, integracao real com servicos de comunicacao '
    'e banco de dados em nuvem.')
add_body(doc,
    'A Neon Cloud garante disponibilidade 24/7, conexoes seguras via SSL e escalabilidade '
    'sem necessidade de administracao de servidor. A integracao Twilio eleva o nivel de '
    'resposta com escalacao real de emergencia com confirmacao humana via teclado telefonico.')

add_heading(doc, 'Resultados Alcancados', level=2)
add_bullet(doc, [
    'Sistema full-stack funcional: API REST + Mobile Expo + PostgreSQL Neon Cloud.',
    'Autenticacao completa com hash bcrypt e recuperacao de conta.',
    'Fluxo SOS de 3 fases com Twilio WhatsApp + Voz + Webhook de resolucao.',
    'Arquitetura MVC em TypeScript com codegen automatico (OpenAPI → Zod + React Query).',
    '9 telas implementadas em React Native com tema escuro premium.',
    '8 tabelas no banco de dados com schema gerenciado pelo Drizzle ORM.',
    'Controle de acesso por genero com redirect automatico.',
    'Logs estruturados (Pino) e validacao de entrada (Zod) em todas as rotas.',
])

add_heading(doc, 'Proximos Passos', level=2)
add_bullet(doc, [
    'Implementar autenticacao JWT para sessoes stateless em producao.',
    'Adicionar notificacoes push (Expo Push Notifications).',
    'Integrar GPS real com rastreamento em tempo real.',
    'Deploy: API no Railway/Render + App nas lojas Google Play e App Store.',
    'Adicionar testes automatizados (Jest + Supertest + Detox).',
])

doc.add_page_break()

# ── ANEXO A – ENDPOINTS ──────────────────────────────────────────────────────
add_heading(doc, 'Anexo A – Endpoints da API REST')
endpoints = [
    ('POST',   '/api/users',                         'Criar usuario'),
    ('GET',    '/api/users/:userId',                 'Buscar usuario'),
    ('PATCH',  '/api/users/:userId',                 'Atualizar usuario'),
    ('POST',   '/api/auth/login',                    'Login'),
    ('POST',   '/api/auth/recover',                  'Recuperar conta'),
    ('POST',   '/api/alerts',                        'Criar alerta SOS + WhatsApp'),
    ('GET',    '/api/alerts/user/:userId',           'Listar alertas'),
    ('PUT',    '/api/alerts/:id/resolve',            'Resolver alerta'),
    ('PUT',    '/api/alerts/:id/webhook/respond',    'Webhook Twilio'),
    ('POST',   '/api/alerts/trigger-voice',          'Disparar ligacao de voz'),
    ('GET',    '/api/contacts/:userId',              'Listar contatos'),
    ('POST',   '/api/contacts',                      'Criar contato'),
    ('DELETE', '/api/contacts/:id',                  'Remover contato'),
    ('GET',    '/api/posts',                         'Listar posts'),
    ('POST',   '/api/posts',                         'Publicar post'),
    ('GET',    '/api/settings/:userId',              'Buscar configuracoes'),
    ('PUT',    '/api/settings/:userId',              'Atualizar configuracoes'),
    ('GET',    '/api/incidents/:userId',             'Listar boletins'),
    ('POST',   '/api/incidents',                     'Registrar boletim'),
    ('GET',    '/api/healthz',                       'Health check'),
]
add_req_table(doc, ['Metodo','Rota','Descricao'], endpoints, [2.5, 8, 7])

# ── ANEXO B – MONOREPO ───────────────────────────────────────────────────────
add_heading(doc, 'Anexo B – Estrutura do Monorepo (pnpm workspaces)')
tree_lines = [
    'SOS-Mulher/  (monorepo pnpm)',
    '├── artifacts/',
    '│   ├── api-server/     # Backend Express 5 + TypeScript',
    '│   │   └── src/ (routes/, controllers/, services/, models/)',
    '│   ├── mobile/         # Expo React Native + Expo Router',
    '│   │   └── app/ (onboarding, (tabs)/, restricted)',
    '│   └── mockup-sandbox/ # Prototipacao de componentes',
    '├── lib/',
    '│   ├── db/             # Schema Drizzle ORM + migracoes',
    '│   ├── api-spec/       # OpenAPI 3.1 + orval.config.ts',
    '│   ├── api-client-react/  # React Query hooks (gerados)',
    '│   └── api-zod/        # Zod schemas (gerados)',
    '├── scripts/',
    '├── pnpm-workspace.yaml',
    '└── package.json',
]
p = doc.add_paragraph()
run = p.add_run('\n'.join(tree_lines))
run.font.name = 'Courier New'
run.font.size = Pt(9)

# save
OUT = '/home/runner/workspace/PI2_EquipeMaria Janaina_Avaliacao2_DocumentacaoFinal.docx'
doc.save(OUT)
print(f'DOCX OK — {os.path.getsize(OUT)//1024} KB')
